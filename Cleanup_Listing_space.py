#!/usr/bin/env python3

import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


def is_empty(paragraph):
    return paragraph.text.strip() == ""


def delete_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)


def insert_empty_paragraph_after(paragraph):
    new_p = OxmlElement("w:p")
    paragraph._element.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def insert_empty_paragraph_before(paragraph):
    new_p = OxmlElement("w:p")
    paragraph._element.addprevious(new_p)
    return Paragraph(new_p, paragraph._parent)


def adjust_style_listing_spacing(doc, style_name):
    """
    For each contiguous block of paragraphs using style_name:
      - Remove an empty paragraph immediately before the block.
      - Ensure there is an empty paragraph immediately after the block.
    """
    paragraphs = doc.paragraphs

    i = 0
    while i < len(paragraphs):

        if paragraphs[i].style.name != style_name:
            i += 1
            continue

        start = i
        while (
            i + 1 < len(paragraphs)
            and paragraphs[i + 1].style.name == style_name
        ):
            i += 1
        end = i

        # Remove blank before block
        if start > 0 and is_empty(paragraphs[start - 1]):
            delete_paragraph(paragraphs[start - 1])
            paragraphs = doc.paragraphs
            start -= 1
            end -= 1

        paragraphs = doc.paragraphs

        # Ensure blank after block
        if end == len(paragraphs) - 1:
            insert_empty_paragraph_after(paragraphs[end])
        elif not is_empty(paragraphs[end + 1]):
            insert_empty_paragraph_after(paragraphs[end])

        paragraphs = doc.paragraphs
        i = end + 1


def adjust_table_spacing(doc):
    """
    Ensure every table has an empty paragraph immediately
    before and after it.
    """
    tables = list(doc.tables)

    for table in tables:
        tbl = table._tbl
        parent = tbl.getparent()

        prev = tbl.getprevious()
        if prev is None or prev.tag.endswith("}tbl"):
            new_p = OxmlElement("w:p")
            parent.insert(parent.index(tbl), new_p)
        elif prev.tag.endswith("}p"):
            para = Paragraph(prev, table._parent)
            if not is_empty(para):
                new_p = OxmlElement("w:p")
                parent.insert(parent.index(tbl), new_p)

        nxt = tbl.getnext()
        if nxt is None or nxt.tag.endswith("}tbl"):
            new_p = OxmlElement("w:p")
            parent.insert(parent.index(tbl) + 1, new_p)
        elif nxt.tag.endswith("}p"):
            para = Paragraph(nxt, table._parent)
            if not is_empty(para):
                new_p = OxmlElement("w:p")
                parent.insert(parent.index(tbl) + 1, new_p)


def main():
    if len(sys.argv) != 2:
        print(f"Usage: python {Path(sys.argv[0]).name} input.docx")
        sys.exit(1)

    input_file = Path(sys.argv[1])

    doc = Document(str(input_file))

    adjust_style_listing_spacing(doc, "AlphabetListing")
    adjust_table_spacing(doc)

    output = input_file.with_name(
        f"{input_file.stem}_fix_spacing{input_file.suffix}"
    )
    doc.save(str(output))

    print(f"Saved: {output}")


if __name__ == "__main__":
    main()